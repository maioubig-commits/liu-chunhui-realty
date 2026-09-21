#!/usr/bin/env python3
"""房仲網站 物件自動更新

讀「買賣案件總表.xlsx」→ 重建 script.js 的 listings → 部署 Netlify → 推 GitHub。
總表沒變就直接結束，所以可以放心定時跑。

用法：
  python3 sync_listings.py            # 有變更才更新並發佈
  python3 sync_listings.py --dry-run  # 只顯示會變什麼，不寫檔、不發佈
  python3 sync_listings.py --no-push  # 更新 script.js 但不部署、不推 GitHub
  python3 sync_listings.py --force    # 總表沒變也強制重跑
"""
import datetime
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zlib

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.dirname(HERE)
XLSX = '/Users/leo/Desktop/1.廣告/買賣案件總表.xlsx'
SCRIPT_JS = os.path.join(SITE, 'script.js')
STATE = os.path.join(HERE, 'state.json')
PHOTO_CACHE = os.path.join(HERE, 'photo_cache.json')
LOG = os.path.join(HERE, 'sync.log')
NETLIFY_SITE = '15b58d8f-8b0f-4d9b-8beb-b0c5c16bc591'
DEPLOY_FILES = ['index.html', 'style.css', 'script.js', 'agent-photo.jpg',
                'line-qr.png', 'og-image.jpg', 'robots.txt', 'sitemap.xml']
# launchd 的 PATH 很精簡，補上 homebrew
os.environ['PATH'] = '/opt/homebrew/bin:/usr/local/bin:' + os.environ.get('PATH', '')

DRY = '--dry-run' in sys.argv
NO_PUSH = '--no-push' in sys.argv
FORCE = '--force' in sys.argv

PLACEHOLDER = open(os.path.join(HERE, 'placeholder.txt')).read().strip()

SHEETS = {'松-電梯': 'apartment', '松-公寓': 'house', '松-店面': 'office',
          '外區-北市': None, '外區-新北市': None, '外區-其他縣市': None}


def log(msg):
    line = f"[{datetime.datetime.now():%Y-%m-%d %H:%M:%S}] {msg}"
    print(line)
    if not DRY:
        with open(LOG, 'a', encoding='utf-8') as f:
            f.write(line + '\n')


def run(cmd, cwd=SITE, check=True):
    r = subprocess.run(cmd, cwd=cwd, capture_output=True, text=True)
    if check and r.returncode != 0:
        raise RuntimeError(f"{' '.join(cmd)} 失敗：{(r.stderr or r.stdout).strip()[-600:]}")
    return r


# ---------- 讀總表 ----------
def load_master():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    rows = []
    for sn, forced in SHEETS.items():
        ws = wb[sn]
        for r in range(3, ws.max_row + 1):
            row = []
            for c in range(1, 16):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, (datetime.date, datetime.datetime)):
                    v = v.isoformat()
                row.append(v)
            if row[3] is None:
                continue
            note = row[14]
            adv = (row[1] == 'O' or row[2] == 'O') and (not note or '不可廣告' not in str(note))
            rows.append({'sheet': sn, 'forced_type': forced, 'title': str(row[3]).strip(),
                         'addr': row[4], 'area': row[5], 'price': row[6], 'layout': row[7],
                         'usage': row[13], 'advertisable': adv})
    return rows


def norm(s):
    s = s.replace('⼭', '山').replace('⺠', '民').replace('⾦', '金').replace('⾯', '面')
    return re.sub(r'[\s_\-（）()【】「」\'"]', '', s)


DISTRICT_CITY = {}
for city, ds in {
    '台北市': '中正 大同 中山 松山 大安 萬華 信義 士林 北投 內湖 南港 文山',
    '新北市': '板橋 三重 中和 永和 新莊 新店 樹林 鶯歌 三峽 淡水 汐止 汐⽌ 瑞芳 土城 蘆洲 五股 泰山 林口 深坑 石碇 坪林 三芝 石門 八里 平溪 雙溪 貢寮 金山 萬里 烏來',
    '桃園市': '桃園 中壢 平鎮 八德 楊梅 蘆竹 大溪 龍潭 龜山 大園 觀音 新屋 復興',
    '新竹市': '東 北 香山',
    '台中市': '中 西 南 北屯 西屯 南屯 太平 大里 霧峰 烏日 豐原',
}.items():
    for d in ds.split():
        DISTRICT_CITY[d + '區'] = city
COUNTIES = ['宜蘭縣', '彰化縣', '南投縣', '雲林縣', '嘉義縣', '屏東縣', '花蓮縣', '台東縣',
            '澎湖縣', '金門縣', '連江縣', '新竹縣', '苗栗縣', '基隆市']


def region_of(sheet, addr):
    if sheet.startswith('松-') or not addr:
        return '台北市'
    addr = str(addr)
    for c in COUNTIES:
        if addr.startswith(c):
            return c
    m = re.match(r'^(\S{2,3}區)', addr)
    if m and m.group(1) in DISTRICT_CITY:
        return DISTRICT_CITY[m.group(1)]
    return '新北市' if sheet == '外區-新北市' else '台北市'


def parse_price(p):
    m = re.search(r'([\d.]+)', str(p)) if p is not None else None
    return int(float(m.group(1))) if m else None


def parse_layout(layout):
    if not layout:
        return 0, 0

    def num(x):
        x = x.strip()
        try:
            return sum(float(p) for p in x.split('+')) if '+' in x else float(x)
        except ValueError:
            m = re.search(r'[\d.]+', x)
            return float(m.group()) if m else 0
    parts = str(layout).strip().split('/')
    rooms = num(parts[0])
    baths = num(parts[2]) if len(parts) >= 3 else (num(parts[-1]) if len(parts) >= 2 else 0)
    return (int(rooms) if rooms == int(rooms) else rooms,
            int(baths) if baths == int(baths) else baths)


CN_NUM = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '七': 7, '八': 8, '九': 9}


def cn_to_int(t):
    if t == '十':
        return 10
    if t.startswith('十'):
        return 10 + CN_NUM.get(t[1:], 0)
    if '十' in t:
        a, _, b = t.partition('十')
        return CN_NUM.get(a, 0) * 10 + CN_NUM.get(b, 0)
    return CN_NUM.get(t, 0)


def parse_floor(addr, title):
    """從地址抓所在樓層（B1 = -1）；抓不到回傳 None，不亂猜。"""
    a = str(addr or '').translate(str.maketrans('０１２３４５６７８９ＢＦ', '0123456789BF'))
    m = re.search(r'B(\d)\s*$|地下\s*(\d)?', a)
    if m:
        return -int(m.group(1) or m.group(2) or 1)
    m = re.search(r'(\d+)\s*(?:[-~至、,]\s*\d+\s*)?(?:樓|F)', a) or re.search(r'([一二三四五六七八九十]+)樓', a)
    if m:
        g = m.group(1)
        return int(g) if g.isdigit() else cn_to_int(g)
    if re.search(r'(?<![\d一二三四五六七八九十])(1樓|一樓)', str(title)):
        return 1
    return None


def classify(r):
    if r['forced_type']:
        return r['forced_type']
    t, u = r['title'], str(r['usage'] or '')
    if any(k in t for k in ['店面', '金店', '店鋪', '收租店', '辦公', '純辦', '住辦']) or \
       any(k in u for k in ['辦公', '事務所', '零售業', '商業用', '店鋪', '市場']):
        return 'office'
    if any(k in t for k in ['透天', '別墅', '公寓', '農舍']):
        return 'house'
    return 'apartment'


# ---------- script.js 讀寫 ----------
ENTRY_RE = re.compile(
    r"\{ id: (\d+), type: '(\w+)', title: '((?:[^'\\]|\\.)*)', region: '([^']*)', price: (\d+), "
    r"area: ([\d.]+), rooms: ([\d.]+), baths: ([\d.]+), (?:floor: (?:-?\d+|null), )?badge: '([^']*)', img: '([^']*)', link: '([^']*)' \}")
ARRAY_RE = re.compile(r"(const listings = \[\n)(.*?)(\n  \];)", re.S)


def current_entries(js):
    out = []
    for m in ENTRY_RE.finditer(js):
        out.append({'type': m.group(2), 'title': m.group(3).replace("\\'", "'"),
                    'price': int(m.group(5)), 'area': float(m.group(6)),
                    'img': m.group(10), 'link': m.group(11), 'badge': m.group(9)})
    return out


def js_str(s):
    return s.replace('\\', '\\\\').replace("'", "\\'")


def fmt_num(n):
    return str(int(n)) if float(n) == int(n) else str(n)


def build(master, cache):
    final, missing = [], 0
    used = {}
    for r in master:
        if not r['advertisable']:
            continue
        if any(k in str(r['layout'] or '') for k in ['純土地', '車位']):
            continue
        key = norm(r['title'])
        area = float(r['area']) if r['area'] not in (None, '') else 0
        cands = cache.get(key, [])
        # 同名多筆時，挑坪數最接近、且還沒被用掉的
        cands = [c for i, c in enumerate(cands) if (key, i) not in used]
        photo = min(cands, key=lambda c: abs(c['area'] - area)) if cands else None
        if photo:
            used[(key, cache[key].index(photo))] = 1
        price = parse_price(r['price'])
        if price is None:
            continue
        rooms, baths = parse_layout(r['layout'])
        e = {'type': photo['type'] if photo else classify(r), 'title': r['title'],
             'region': region_of(r['sheet'], r['addr']), 'price': price, 'area': area,
             'rooms': rooms, 'baths': baths, 'floor': parse_floor(r['addr'], r['title'])}
        if photo:
            e.update(img=photo['img'], link=photo['link'], badge=photo['badge'])
        else:
            e.update(img=PLACEHOLDER, link='', badge='照片準備中')
            missing += 1
        final.append(e)
    final.sort(key=lambda e: (e['price'], e['title']))
    return final, missing


def render(final):
    lines = []
    for i, e in enumerate(final, 1):
        lines.append(
            f"    {{ id: {i}, type: '{e['type']}', title: '{js_str(e['title'])}', region: '{e['region']}', "
            f"price: {e['price']}, area: {fmt_num(e['area'])}, rooms: {fmt_num(e['rooms'])}, "
            f"baths: {fmt_num(e['baths'])}, floor: {'null' if e['floor'] is None else e['floor']}, badge: '{e['badge']}', img: '{e['img']}', link: '{e['link']}' }},")
    return '\n'.join(lines)


def main():
    if not os.path.exists(XLSX):
        log(f'找不到總表：{XLSX}')
        return 1
    xlsx_hash = hashlib.md5(open(XLSX, 'rb').read()).hexdigest()
    state = json.load(open(STATE)) if os.path.exists(STATE) else {}
    if not FORCE and state.get('xlsx_hash') == xlsx_hash and state.get('published'):
        return 0  # 沒變，安靜結束

    js = open(SCRIPT_JS, encoding='utf-8').read()
    old = current_entries(js)
    if not old:
        log('script.js 讀不到物件，中止')
        return 1

    # 照片快取：以目前網站上有真照片的為底，累積起來（下架再回來也找得回照片）
    cache = json.load(open(PHOTO_CACHE, encoding='utf-8')) if os.path.exists(PHOTO_CACHE) else {}
    for e in old:
        if not e['img'].startswith('data:'):
            lst = cache.setdefault(norm(e['title']), [])
            if not any(c['img'] == e['img'] for c in lst):
                lst.append({k: e[k] for k in ('type', 'area', 'img', 'link', 'badge')})

    try:
        master = load_master()
    except Exception as ex:  # Excel 存檔到一半等情況，下次再來
        log(f'總表讀取失敗，稍後重試：{ex}')
        return 1
    final, missing = build(master, cache)

    # 安全閥：筆數暴增暴減多半是總表格式被動過，不要直接上線
    if len(final) < len(old) * 0.7 or len(final) > len(old) * 1.4:
        log(f'物件數從 {len(old)} 變成 {len(final)}，變動過大，已中止（需人工確認）')
        return 2

    old_keys = {(norm(e['title']), e['price']) for e in old}
    new_keys = {(norm(e['title']), e['price']) for e in final}
    added = [e['title'] for e in final if (norm(e['title']), e['price']) not in old_keys]
    removed = [e['title'] for e in old if (norm(e['title']), e['price']) not in new_keys]
    summary = f'{len(old)} → {len(final)} 筆（新增 {len(added)}、下架/調價 {len(removed)}、待補照片 {missing}）'
    log(summary)
    for t in added[:15]:
        log(f'  + {t}')
    for t in removed[:15]:
        log(f'  - {t}')
    if DRY:
        return 0

    new_js = ARRAY_RE.sub(lambda m: m.group(1) + render(final) + m.group(3), js, count=1)
    with tempfile.NamedTemporaryFile('w', suffix='.js', delete=False, encoding='utf-8') as f:
        f.write(new_js)
    try:
        run(['node', '--check', f.name])
    finally:
        os.unlink(f.name)

    open(SCRIPT_JS, 'w', encoding='utf-8').write(new_js)
    json.dump(cache, open(PHOTO_CACHE, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    state.update(xlsx_hash=xlsx_hash, published=False, count=len(final))
    json.dump(state, open(STATE, 'w'), indent=1)

    if NO_PUSH:
        log('已更新 script.js（--no-push，未部署）')
        return 0

    # 部署乾淨包（不直接傳整個資料夾，避免雜檔上線）
    pkg = tempfile.mkdtemp(prefix='chunhui-deploy-')
    try:
        for name in DEPLOY_FILES:
            shutil.copy(os.path.join(SITE, name), pkg)
        r = run(['netlify', 'deploy', '--prod', '--dir', pkg, '--site', NETLIFY_SITE,
                 '--message', f'自動同步總表 {summary}'])
        log('Netlify 部署完成')
    finally:
        shutil.rmtree(pkg, ignore_errors=True)

    run(['git', 'add', 'script.js', 'tools/photo_cache.json'])
    if run(['git', 'diff', '--cached', '--quiet'], check=False).returncode != 0:
        run(['git', 'commit', '-m', f'自動同步買賣案件總表：{summary}\n\n'
             'Co-Authored-By: Claude Sonnet 5 <noreply@anthropic.com>'])
        try:
            run(['git', 'push'])
            log('GitHub 已備份')
        except RuntimeError as ex:
            log(f'GitHub 推送失敗（網站已更新，不影響上線）：{ex}')
    state['published'] = True
    json.dump(state, open(STATE, 'w'), indent=1)
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception as ex:
        log(f'發生錯誤：{ex}')
        sys.exit(1)
